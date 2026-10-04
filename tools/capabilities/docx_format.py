"""Convert plain text / light markdown into styled Word paragraphs."""
from __future__ import annotations

import os
import re
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    import docx.document

_H1 = re.compile(r"^#\s+(.+)$")
_H2 = re.compile(r"^##\s+(.+)$")
_H3 = re.compile(r"^###\s+(.+)$")
_BOLD_LINE = re.compile(r"^\*\*(.+?)\*\*\s*$")
_BULLET = re.compile(r"^[\-\*•]\s+(.+)$")
_NUMBERED = re.compile(r"^(\d+)[\.\)]\s+(.+)$")
_INLINE = re.compile(r"(\*\*(.+?)\*\*|\*(.+?)\*|__(.+?)__|_(.+?)_)")

_KNOWN_SECTIONS = frozenset(
    {
        "introduction",
        "conclusion",
        "abstract",
        "references",
        "executive summary",
        "body section i",
        "body section ii",
        "body section 1",
        "body section 2",
    }
)


def _strip_md_heading_markers(text: str) -> str:
    text = text.strip()
    for pattern in (_H1, _H2, _H3):
        match = pattern.match(text)
        if match:
            return match.group(1).strip()
    match = _BOLD_LINE.match(text)
    if match:
        return match.group(1).strip()
    return text


def _add_inline_runs(paragraph, text: str) -> None:
    """Add runs to a paragraph, interpreting **bold** and *italic*."""
    text = text.strip()
    if not text:
        return
    pos = 0
    for match in _INLINE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        if match.group(2):
            run = paragraph.add_run(match.group(2))
            run.bold = True
        elif match.group(3):
            run = paragraph.add_run(match.group(3))
            run.italic = True
        elif match.group(4):
            run = paragraph.add_run(match.group(4))
            run.bold = True
        elif match.group(5):
            run = paragraph.add_run(match.group(5))
            run.italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _is_section_title(line: str) -> bool:
    stripped = line.strip()
    if not stripped or len(stripped) > 72 or stripped.startswith("http"):
        return False
    if "**" in stripped or stripped.startswith("#"):
        return False
    if stripped.lower().rstrip(":") in _KNOWN_SECTIONS:
        return True
    if stripped.endswith(".") and not stripped.endswith("..."):
        return False
    words = stripped.split()
    if not 1 <= len(words) <= 8:
        return False
    return all(word[0].isupper() for word in words if word and word[0].isalpha())


def _heading_level(line: str) -> int | None:
    stripped = line.strip()
    if _H1.match(stripped):
        return 1
    if _H2.match(stripped):
        return 2
    if _H3.match(stripped):
        return 3
    if _BOLD_LINE.match(stripped):
        return 2
    if stripped.endswith(":") and len(stripped) < 80 and not stripped.startswith("http"):
        if stripped[0].isupper() and "**" not in stripped:
            return 3
    if _is_section_title(stripped):
        return 2
    return None


def populate_document(document: "docx.document.Document", content: str) -> None:
    """Fill an empty python-docx Document from assignment-style text."""
    lines = str(content).replace("\r\n", "\n").split("\n")
    title_used = False
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue

        stripped = line.strip()
        if not title_used and len(stripped.split()) <= 14 and _heading_level(stripped) is None:
            if not stripped.startswith(("-", "*", "http")) and not _NUMBERED.match(stripped):
                document.add_heading(_strip_md_heading_markers(stripped), level=0)
                title_used = True
                continue

        level = _heading_level(line)
        if level is not None:
            title = _strip_md_heading_markers(line)
            document.add_heading(title, level=level)
            continue

        bullet = _BULLET.match(line.strip())
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_runs(paragraph, bullet.group(1))
            continue

        numbered = _NUMBERED.match(line.strip())
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            _add_inline_runs(paragraph, numbered.group(2))
            continue

        paragraph = document.add_paragraph()
        _add_inline_runs(paragraph, line.strip())


def write_formatted_docx(path: str, content: str) -> tuple[bool, str]:
    try:
        import docx
    except ImportError:
        return False, "DOCX create requires: pip install python-docx"
    try:
        parent = os.path.dirname(path)
        if parent:
            os.makedirs(parent, exist_ok=True)
        document = docx.Document()
        populate_document(document, content)
        document.save(path)
        return True, f"Document saved to {path}."
    except Exception as exc:
        return False, f"Could not write DOCX: {exc}"
